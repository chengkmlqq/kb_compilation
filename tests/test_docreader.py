"""docreader migration tests — registry + parser integration (migrated from WeKnora)."""

from __future__ import annotations

import io

import pytest

from docreader.models.document import Chunk, Document
from docreader.parser import Parser
from docreader.parser.registry import BUILTIN_ENGINE, registry
from docreader.splitter.splitter import TextSplitter


def test_document_model_roundtrip() -> None:
    doc = Document(content="hello", metadata={"k": "v"})
    assert doc.is_valid()
    assert doc.get_content() == "hello"
    chunk = Chunk(content="c", seq=0, start=0, end=1)
    assert Chunk.from_dict(chunk.to_dict()).content == "c"


def test_registry_builtin_covers_formats() -> None:
    engines = registry.get_engine_names()
    assert BUILTIN_ENGINE in engines
    assert "markitdown" in engines
    for ft in ("md", "pdf", "docx", "xlsx", "epub", "mhtml"):
        cls = registry.get_parser_class(BUILTIN_ENGINE, ft)
        assert cls is not None


def test_registry_fallback_to_builtin() -> None:
    # Requesting markitdown engine for a type it does not support must
    # gracefully fall back to the builtin engine rather than raise.
    cls = registry.get_parser_class("markitdown", "epub")
    assert cls.__name__ == "EPUBParser"


def test_registry_unsupported_type_raises() -> None:
    with pytest.raises(ValueError):
        registry.get_parser_class(BUILTIN_ENGINE, "not-a-real-type")


def test_markdown_parse() -> None:
    parser = Parser()
    md = b"# Title\n\nIntro paragraph.\n\n## Section\n\nSection body text.\n"
    doc = parser.parse_file("t.md", "md", md)
    assert "# Title" in doc.content
    assert "Section body text." in doc.content


def test_markitdown_engine_parses_xlsx() -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "name"
    ws["B1"] = "age"
    ws["A2"] = "alice"
    ws["B2"] = 30
    buf = io.BytesIO()
    wb.save(buf)

    doc = Parser().parse_file("t.xlsx", "xlsx", buf.getvalue(), parser_engine="markitdown")
    assert "alice" in doc.content
    assert "| name | age |" in doc.content


def test_text_splitter_splits_and_restores() -> None:
    text = "This is a sentence. " * 60  # long enough to force a split
    splitter = TextSplitter(chunk_size=200, chunk_overlap=20)
    chunks = splitter.split_text(text)
    assert len(chunks) > 1
    # restore_text reassembles (within overlap tolerance)
    restored = splitter.restore_text(chunks)
    assert "This is a sentence." in restored
